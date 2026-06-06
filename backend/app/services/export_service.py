import io
import json
import zipfile
from typing import Dict, Any
from fpdf import FPDF
from pptx import Presentation
from pptx.util import Inches, Pt

def get_project_data(db, project_id: str, version_number: int = None) -> Dict[str, Any]:
    from app.db import queries
    project = queries.get_project_by_id(db, project_id)
    if not project:
        raise ValueError("Project not found")

    if version_number:
        res = db.table("project_versions").select("*").eq("project_id", project_id).eq("version_number", version_number).execute()
        if not res.data:
            raise ValueError(f"Version {version_number} not found")
        outputs = json.loads(res.data[0].get("summary") or "{}")
        critic = outputs.get("critic", {})
        score = res.data[0].get("critic_score") or critic.get("overall_score")
        return {
            "title": project["title"],
            "description": project["description"],
            "outputs": outputs,
            "version": version_number,
            "score": score
        }
    else:
        runs = queries.get_agent_runs_for_project(db, project_id)
        outputs = {r["agent_name"]: r.get("output") for r in runs}
        critic = outputs.get("critic", {})
        score = critic.get("overall_score", 0)
        return {
            "title": project["title"],
            "description": project["description"],
            "outputs": outputs,
            "version": "Latest",
            "score": score
        }

def export_json(db, project_id: str, version_number: int = None) -> bytes:
    data = get_project_data(db, project_id, version_number)
    return json.dumps(data, indent=2).encode('utf-8')

def export_markdown(db, project_id: str, version_number: int = None) -> bytes:
    data = get_project_data(db, project_id, version_number)
    outputs = data["outputs"]
    
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        readme = f"# {data['title']}\n\n{data['description']}\n\nVersion: {data['version']}\nScore: {data['score']}\n"
        zf.writestr("README.md", readme)
        
        if outputs.get("planner"):
            zf.writestr("planner.md", f"# Planner Output\n\n```json\n{json.dumps(outputs['planner'], indent=2)}\n```")
        if outputs.get("pm"):
            zf.writestr("pm.md", f"# PM Output\n\n```json\n{json.dumps(outputs['pm'], indent=2)}\n```")
        if outputs.get("architect"):
            zf.writestr("architecture.md", f"# Architecture Output\n\n```json\n{json.dumps(outputs['architect'], indent=2)}\n```")
        if outputs.get("database"):
            zf.writestr("database.md", f"# Database Output\n\n```json\n{json.dumps(outputs['database'], indent=2)}\n```")
        if outputs.get("documentation"):
            docs = outputs["documentation"]
            zf.writestr("documentation.md", f"# Documentation\n\n## README\n{docs.get('readme', '')}\n\n## API Docs\n{docs.get('api_docs', '')}\n\n## Setup\n{docs.get('setup_guide', '')}")
        if outputs.get("critic"):
            zf.writestr("critic_review.md", f"# Critic Review\n\nScore: {data['score']}\n\n```json\n{json.dumps(outputs['critic'], indent=2)}\n```")
            
    return zip_buffer.getvalue()

def export_pdf(db, project_id: str, version_number: int = None) -> bytes:
    data = get_project_data(db, project_id, version_number)
    outputs = data["outputs"]
    
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", 'B', 24)
    # fpdf2 allows multi_cell string handling without extreme error assuming latin-1
    # We'll encode it safely just in case
    title_safe = data['title'].encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(w=0, h=10, text=title_safe, align='L', new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", '', 12)
    pdf.multi_cell(w=0, h=10, text=f"Version: {data['version']} | Quality Score: {data['score']}", align='L', new_x="LMARGIN", new_y="NEXT")
    pdf.ln(10)
    
    pdf.set_font("Helvetica", 'B', 16)
    pdf.cell(0, 10, "Executive Summary", align='L', new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", '', 12)
    desc_safe = data['description'].encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(w=0, h=8, text=desc_safe, align='L', new_x="LMARGIN", new_y="NEXT")
    pdf.ln(10)

    def add_section(title, content_dict):
        if not content_dict: return
        pdf.add_page()
        pdf.set_font("Helvetica", 'B', 16)
        pdf.cell(0, 10, title, align='L', new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", '', 10)
        for k, v in content_dict.items():
            pdf.set_font("Helvetica", 'B', 10)
            pdf.cell(0, 8, str(k).replace('_', ' ').title(), align='L', new_x="LMARGIN", new_y="NEXT")
            pdf.set_font("Helvetica", '', 10)
            if isinstance(v, list):
                for item in v:
                    text = f"- {item}".encode('latin-1', 'replace').decode('latin-1')
                    pdf.multi_cell(w=0, h=6, text=text, align='L', new_x="LMARGIN", new_y="NEXT")
            else:
                text = str(v).encode('latin-1', 'replace').decode('latin-1')
                pdf.multi_cell(w=0, h=6, text=text, align='L', new_x="LMARGIN", new_y="NEXT")
            pdf.ln(4)

    add_section("Requirements (PM)", outputs.get("pm"))
    add_section("Architecture", outputs.get("architect"))
    add_section("Database Design", outputs.get("database"))
    add_section("Critic Review", outputs.get("critic"))
    
    return bytes(pdf.output())

def export_pptx(db, project_id: str, version_number: int = None) -> bytes:
    data = get_project_data(db, project_id, version_number)
    outputs = data["outputs"]
    
    prs = Presentation()
    
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = data["title"]
    subtitle.text = f"Architecture Review (Version: {data['version']} | Score: {data['score']})"
    
    def add_slide(title_text, content_text):
        bullet_slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        title_shape.text = title_text
        tf = body_shape.text_frame
        tf.text = content_text[:2000] + ("..." if len(content_text) > 2000 else "")
        
    add_slide("Project Overview", data["description"])
    
    if outputs.get("pm"):
        features = outputs["pm"].get("core_features", [])
        add_slide("Core Features", "\n".join([f"• {f}" for f in features]))
        
    if outputs.get("architect"):
        components = outputs["architect"].get("components", [])
        add_slide("Architecture Components", "\n".join([f"• {c.get('name', c)}" if isinstance(c, dict) else f"• {c}" for c in components]))
        
    if outputs.get("database"):
        tables = outputs["database"].get("tables", [])
        add_slide("Database Entities", "\n".join([f"• {t.get('name', t)}" if isinstance(t, dict) else f"• {t}" for t in tables]))
        
    if outputs.get("critic"):
        weaknesses = outputs["critic"].get("weaknesses", [])
        add_slide("Critic Findings", "\n".join([f"• {w}" for w in weaknesses]))
        
    pptx_io = io.BytesIO()
    prs.save(pptx_io)
    return pptx_io.getvalue()
